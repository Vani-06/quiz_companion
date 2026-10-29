"""Builds the Phase 3 final PPT from the Phase 2 deck (same theme, fonts and layout).

Usage: python docs/build_phase3_deck.py <phase2.pptx> <shots_dir> <output.pptx>
"""
import copy
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

src, shots, out = sys.argv[1:4]
NAVY, GOLD, CODE_BG, GREY, LIGHT = "14284B", "C9962B", "0F1B33", "4A5563", "F3F4F6"
prs = Presentation(src)
layout = prs.slides[16].slide_layout          # the content layout used by Phase 2 slides


def rgb(h):
    return RGBColor.from_string(h)


# ---- keep only the title slide, retitle it -------------------------------------
ids = prs.slides._sldIdLst
for sld in list(ids)[1:]:
    prs.part.drop_rel(sld.rId)
    ids.remove(sld)
for sh in prs.slides[0].shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() == "Adaptive Learning Platform":
                    r.text = "SkillSync: Adaptive Quiz Recommendation & Learning Analytics Platform"
                    r.font.size = Pt(23)
                r.text = (r.text.replace("Phase 2: Implementation Core", "Phase 3: Integration & Delivery")
                          .replace("Database implementation, data population and business logic",
                                   "Front end, database connectivity, testing and demonstration"))

page = [1]


def text(slide, x, y, w, h, lines, size=14, color="1F2933", bold=False, font="Arial", fill=None,
         anchor=MSO_ANCHOR.TOP, shape=None, pad=0.12, center=False):
    box = slide.shapes.add_shape(shape or MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.line.fill.background()
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(fill)
    else:
        box.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        t, o = (line if isinstance(line, tuple) else (line, {}))
        r = p.add_run()
        r.text = t
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.name = o.get("font", font)
        r.font.color.rgb = rgb(o.get("color", color))
        p.space_after = Pt(o.get("after", 4))
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    return box


def slide(kicker, title):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    page[0] += 1
    text(s, 0.667, 0.42, 12, 0.3, [(kicker, {"bold": True, "color": GOLD, "size": 12})], pad=0)
    text(s, 0.667, 0.78, 12, 0.5, [(title, {"bold": True, "color": NAVY, "size": 27, "font": "Georgia"})], pad=0)
    text(s, 0.667, 6.96, 12, 0.3,
         [(f"SkillSync · Phase 3 · {page[0]}", {"size": 12, "color": "6A7179"})], pad=0)
    return s


def picture(s, name, x, y, w, h):
    """Fit an image inside the (w, h) box, centred."""
    path = f"{shots}/{name}"
    with Image.open(path) as im:
        aspect = im.size[0] / im.size[1]
    pw, ph = (w, w / aspect) if w / h < aspect else (h * aspect, h)
    s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def table(s, x, y, w, widths, header, rows, size=12, row_h=0.36):
    t = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w),
                           Inches(row_h * (len(rows) + 1))).table
    total = sum(widths)
    for i, cw in enumerate(widths):
        t.columns[i].width = Inches(w * cw / total)
    for r, row in enumerate([header] + rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = val
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size, run.font.name, run.font.bold = Pt(size), "Arial", r == 0
            run.font.color.rgb = rgb("000000" if r == 0 else "1F2933")
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("E2E2DC" if r == 0 else "FFFFFF")
            cell.margin_top = cell.margin_bottom = Inches(0.04)
    return t


def code(s, x, y, w, h, lines, size=11):
    return text(s, x, y, w, h, [(l if l else " ", {"font": "Consolas", "size": size, "color": "E6EAF2", "after": 0})
                                for l in lines], fill=CODE_BG, pad=0.18, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def steps(s, y, labels):
    n = len(labels)
    w, gap = 2.05, 0.45
    x0 = 0.667 + (12.0 - (n * w + (n - 1) * gap)) / 2
    for i, lab in enumerate(labels):
        b = text(s, x0 + i * (w + gap), y, w, 0.85,
                 [(f"STEP {i + 1}", {"size": 11, "bold": True, "color": GOLD, "after": 0}),
                  (lab, {"size": 13, "bold": True, "color": "FFFFFF"})],
                 fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE, center=True)
        if i < n - 1:
            text(s, x0 + i * (w + gap) + w, y + 0.22, gap, 0.4, [("→", {"size": 18, "color": GOLD})],
                 center=True, pad=0)


# ---- 2. recap and what Phase 3 adds -------------------------------------------
s = slide("Recap", "From Phase 2 to Phase 3")
table(s, 0.667, 1.45, 12, [2, 5, 5], ["Phase", "Delivered", "Evidence"], [
    ["Phase 1", "EER model, relational schema, normalisation to 3NF", "10 tables, 12 foreign keys"],
    ["Phase 2", "PostgreSQL DDL, sample data, business logic", "36 rows; functions, triggers, views"],
    ["Phase 3", "Web front end, JDBC client, full CRUD, validation, tests, recommendations",
     "7 screens, 31 automated tests, pgvector"],
], size=13, row_h=0.5)
text(s, 0.667, 3.7, 12, 2.4, [
    ("Phase 3 adds", {"bold": True, "size": 16, "color": NAVY}),
    ("A React web app and an Express API that read and write PostgreSQL on every screen.", {}),
    ("Java JDBC and PowerShell ODBC clients that run the same retrieval, DML and business-logic functions.", {}),
    ("Quiz taking graded inside the database, student sign-up, leaderboard and full edit/delete.", {}),
    ("Similar-item recommendations with vector embeddings (pgvector), on top of weak-topic rules.", {}),
], fill=LIGHT, pad=0.2)

# ---- 3. architecture ---------------------------------------------------------
s = slide("Part 9 · Database connectivity", "Connectivity Architecture")
picture(s, "diagram_arch.png", 0.667, 1.4, 8.0, 5.4)
text(s, 8.9, 1.5, 3.8, 5.2, [
    ("Three tiers", {"bold": True, "size": 16, "color": NAVY}),
    ("React screens call an Express API over HTTP.", {}),
    ("The API validates input and runs parameterised SQL through a node-postgres connection pool.", {}),
    ("A Java client (JDBC) and a PowerShell client (ODBC) connect straight to PostgreSQL.", {}),
    ("All rules live in the database, so every client behaves the same way.", {}),
], fill=LIGHT, pad=0.18, size=13)

# ---- 4. CRUD from the front end ----------------------------------------------
s = slide("Part 9 · Database connectivity", "CRUD from the Front End")
table(s, 0.667, 1.4, 12, [1.2, 3, 7.8], ["Operation", "Screen / button", "SQL executed (parameterised)"], [
    ["Create", "Topics → Add topic", "INSERT INTO topic (name, description, parent_topic_id) VALUES ($1, $2, $3)"],
    ["Create", "Learning Items → Add", "INSERT INTO learning_item … + quiz or project + project_tech (one transaction)"],
    ["Read", "Dashboard", "SELECT … FROM learning_record, topic_proficiency, recommendation WHERE user_id = $1"],
    ["Update", "Topics → Edit", "UPDATE topic SET name = $2, description = $3, parent_topic_id = $4 WHERE topic_id = $1"],
    ["Update", "Dashboard → Done / Dismiss", "UPDATE recommendation SET status = $3 WHERE user_id = $1 AND item_id = $2"],
    ["Delete", "Topics → Delete", "DELETE FROM topic WHERE topic_id = $1  (guarded by a trigger)"],
    ["Function", "Quizzes → Submit", "SELECT grade_quiz($1, $2::jsonb); SELECT record_attempt($1, $2, 'completed', $3)"],
    ["Function", "Account → Sign up", "SELECT register_user($1, $2, $3)  (password already hashed)"],
], size=11, row_h=0.5)

# ---- 5. front-end screens ----------------------------------------------------
s = slide("Part 7 · Front end", "Front-End Screens on Live Data")
picture(s, "user_dashboard.png", 0.667, 1.4, 5.9, 3.9)
picture(s, "user_take_quiz.png", 6.77, 1.4, 5.9, 3.9)
text(s, 0.667, 5.35, 5.9, 0.3, [("Dashboard: attempts, proficiency chart, recent attempts", {"size": 11, "color": GREY})], pad=0)
text(s, 6.77, 5.35, 5.9, 0.3, [("Take quiz: graded by grade_quiz() in the database", {"size": 11, "color": GREY})], pad=0)
text(s, 0.667, 5.85, 12, 0.9, [
    ("Seven screens: Dashboard, My Learning, Quizzes, Leaderboard, Learning Items, Topics, Account. "
     "Nothing is hard-coded: every number comes from the 11 PostgreSQL tables.", {"size": 13})],
    fill=LIGHT, pad=0.18, anchor=MSO_ANCHOR.MIDDLE)

# ---- 6. quiz flow -------------------------------------------------------------
s = slide("Part 8 · Business logic", "What Happens When a Student Takes a Quiz")
steps(s, 1.5, ["Submit answers", "grade_quiz scores them", "record_attempt saves it", "Trigger updates proficiency",
               "Result shown"])
table(s, 0.667, 2.8, 12, [2.2, 6, 3.8], ["Rule", "What it does", "Enforced by"], [
    ["Valid scores and dates", "Score 0–100; status from a fixed list; no finish before start", "CHECK constraints"],
    ["Attempt numbering", "Next attempt number is calculated, never typed; advisory lock prevents clashes", "record_attempt()"],
    ["Grading", "Marks earned ÷ total marks × 100; pass if ≥ pass_mark", "grade_quiz()"],
    ["Proficiency", "Latest scored attempt per topic; attempts counted", "Trigger on learning_record"],
    ["Leaderboard", "Only each student's best score per item counts", "View v_leaderboard"],
    ["Protected topics", "A topic that still has items cannot be deleted", "Trigger on topic"],
], size=11, row_h=0.5)

# ---- 7. recommendations --------------------------------------------------------
s = slide("Part 8 · Business logic", "Recommendations: Weak Topics and Vector Similarity")
picture(s, "user_recommendations.png", 0.667, 1.4, 6.6, 5.4)
text(s, 7.5, 1.5, 5.2, 5.2, [
    ("Two methods in one button", {"bold": True, "size": 16, "color": NAVY}),
    ("Topic: every item in a topic below 70, weakest first.", {}),
    ("Similar: each item is a 384-number vector (all-MiniLM-L6-v2). A student's profile is the average of "
     "their attempts, weighted towards low scores.", {}),
    ("Unattempted items closest to the profile are recommended, with a reason such as "
     "“Builds on Stack Operations, which you passed”.", {}),
    ("pgvector keeps embeddings in PostgreSQL, so one SQL query joins them with attempts.", {}),
], fill=LIGHT, pad=0.18, size=13)

# ---- 8. create / update / delete evidence -------------------------------------
s = slide("Part 9 · Database connectivity", "Evidence: Create, Update and Delete")
for i, (img, cap) in enumerate([("user_sql_create.png", "Create: row inserted from the Topics screen"),
                                ("user_sql_update.png", "Update: the new description is saved"),
                                ("user_sql_delete.png", "Delete: no matching row remains")]):
    y = 1.4 + i * 1.5
    picture(s, img, 0.667, y, 7.8, 1.35)
    text(s, 8.7, y + 0.35, 4.0, 0.6, [(cap, {"size": 13, "color": NAVY, "bold": True})], pad=0)
text(s, 0.667, 5.95, 12, 0.85, [
    ("Each action was done in the running web app and checked with a SELECT in psql. The same operations run from "
     "Java through JDBC PreparedStatements; multi-table changes use one transaction.", {"size": 13})],
    fill=LIGHT, pad=0.2, anchor=MSO_ANCHOR.MIDDLE)

# ---- 9. errors ---------------------------------------------------------------------
s = slide("Part 10 · Testing and validation", "Constraint Violations and Error Handling")
picture(s, "user_error_duplicate.png", 0.667, 1.4, 6.0, 3.4)
picture(s, "user_error_trigger.png", 6.77, 1.4, 5.9, 3.4)
text(s, 0.667, 4.85, 6.0, 0.3, [("UNIQUE violation: duplicate topic name", {"size": 11, "color": GREY})], pad=0)
text(s, 6.77, 4.85, 5.9, 0.3, [("Trigger: topic still has learning items", {"size": 11, "color": GREY})], pad=0)
table(s, 0.667, 5.3, 12, [1.5, 3, 1, 6], ["SQLSTATE", "Meaning", "HTTP", "Message"], [
    ["23505", "Unique violation", "409", "That value already exists (unique constraint)."],
    ["23503", "Foreign key violation", "409", "Blocked by a related record (foreign key constraint)."],
    ["23514", "Check violation", "400", "A value is outside the allowed range (check constraint)."],
], size=11, row_h=0.36)

# ---- 10. JDBC ----------------------------------------------------------------------
s = slide("Part 9 · Database connectivity", "JDBC Client")
code(s, 0.667, 1.4, 7.4, 4.3, [
    "String sql = \"INSERT INTO topic (name, description,\"",
    "           + \" parent_topic_id) VALUES (?, ?, ?)\"",
    "           + \" RETURNING topic_id\";",
    "try (PreparedStatement st = db.prepareStatement(sql)) {",
    "    st.setString(1, name);",
    "    st.setString(2, description);",
    "    st.setObject(3, parentId);",
    "    ResultSet rs = st.executeQuery();",
    "}",
    "",
    "// grade and record in one transaction",
    "db.setAutoCommit(false);",
    "SELECT grade_quiz(?, ?::jsonb)",
    "SELECT record_attempt(?, ?, 'completed', ?::numeric)",
    "db.commit();   // db.rollback() on error",
], size=12)
text(s, 8.3, 1.4, 4.4, 4.3, [
    ("jdbc/JdbcDemo.java", {"bold": True, "size": 16, "color": NAVY}),
    ("PostgreSQL JDBC driver 42.7.4, Java 21.", {}),
    ("Menu: list, add, rename and delete topics; list items; take a quiz; leaderboard.", {}),
    ("--demo runs every operation in turn, including a rejected delete.", {}),
    ("Password comes from DB_PASSWORD or a prompt, never from source code.", {}),
], fill=LIGHT, pad=0.18, size=13)

# ---- 10b. ODBC ------------------------------------------------------------------
s = slide("Part 9 · Database connectivity", "ODBC Client")
code(s, 0.667, 1.4, 7.4, 4.3, [
    "$driver = Get-OdbcDriver -Platform '64-bit' |",
    "    Where-Object Name -like '*UNICODE*'",
    "$cs = \"Driver={$($driver.Name)};Server=localhost;\"",
    "    + \"Port=5432;Database=quiz_companion;Uid=postgres;Pwd=...\"",
    "$db = New-Object System.Data.Odbc.OdbcConnection($cs)",
    "$db.Open()",
    "",
    "# parameterised INSERT",
    "INSERT INTO topic (name, description, parent_topic_id)",
    "VALUES (?, ?, ?) RETURNING topic_id",
    "",
    "# grade and record in one transaction",
    "SELECT grade_quiz(?, CAST(? AS jsonb))",
    "SELECT record_attempt(?, ?, 'completed', CAST(? AS numeric))",
    "$tx.Commit()   # $tx.Rollback() on error",
], size=12)
text(s, 8.3, 1.4, 4.4, 4.3, [
    ("odbc/OdbcDemo.ps1", {"bold": True, "size": 16, "color": NAVY}),
    ("PostgreSQL ODBC driver (psqlODBC, Unicode, 64-bit) through System.Data.Odbc.", {}),
    ("Same menu and -Demo run as the JDBC client.", {}),
    ("The demo ends with a rejected delete: trigger error 23001 “Topic 3 still has learning items”.", {}),
    ("The quiz attempt is rolled back, so the sample data stays unchanged.", {}),
], fill=LIGHT, pad=0.18, size=13)

# ---- 11. testing -------------------------------------------------------------------
s = slide("Part 10 · Testing and validation", "Test Strategy and Results")
table(s, 0.667, 1.4, 6.4, [3, 1, 1, 1], ["Category", "Total", "Passed", "Failed"], [
    ["Functional", "9", "9", "0"],
    ["Constraint violation (incl. triggers, functions)", "10", "10", "0"],
    ["Error handling", "4", "4", "0"],
    ["Total documented cases", "23", "23", "0"],
], size=13, row_h=0.5)
text(s, 0.667, 4.1, 6.4, 2.6, [
    ("Three levels", {"bold": True, "size": 15, "color": NAVY}),
    ("SQL: business_logic_tests.sql, rolled back.", {}),
    ("API: 31 automated tests on a separate, rebuilt test database.", {}),
    ("Front end, JDBC and ODBC: each run with valid and invalid input.", {}),
    ("One case failed first (NOW() vs TIMESTAMPTZ) and passed after the fix.", {}),
], fill=LIGHT, pad=0.18, size=13)
code(s, 7.3, 1.4, 5.4, 5.3, [
    "PASS connectivity and retrieval",
    "PASS topics CRUD",
    "PASS learning items CRUD",
    "PASS questions CRUD",
    "PASS business rules through the API",
    "PASS accounts and users",
    "PASS error handling",
    "PASS vector recommendations",
    "",
    "tests 31",
    "pass 31",
    "fail 0",
], size=14)

# ---- 12. scope achieved ----------------------------------------------------------
s = slide("Results", "Scope Achieved Against Phase 1")
table(s, 0.667, 1.4, 12, [4, 1.6, 6.4], ["Module (Phase 1 scope)", "Status", "Remarks"], [
    ["Topics, quizzes, projects, questions", "Completed", "Full CRUD from the front end"],
    ["Quiz taking and grading", "Completed", "grade_quiz() and record_attempt() in one transaction"],
    ["Attempt history and topic proficiency", "Completed", "Maintained by trigger trg_update_proficiency"],
    ["Recommendations", "Completed", "Weak-topic rule plus pgvector similarity (replaces Qdrant)"],
    ["Leaderboard", "Completed", "View v_leaderboard"],
    ["Student accounts", "Partly", "Sign-up and sign-in with hashed passwords; no sessions or roles yet"],
    ["Caching of frequent reads (Redis)", "Not built", "Indexes added instead; not needed at this data size"],
    ["Database connectivity", "Completed", "node-postgres for the web app; JDBC and ODBC console clients"],
], size=12, row_h=0.5)

# ---- 13. demo flow ------------------------------------------------------------------
s = slide("Demonstration", "Live Demo Flow")
steps(s, 1.5, ["Sign up", "Take a quiz", "See proficiency", "Refresh recommendations", "Leaderboard"])
text(s, 0.667, 2.9, 12, 3.6, [
    ("Then show the database side", {"bold": True, "size": 16, "color": NAVY}),
    ("psql: SELECT the new attempt and the updated topic_proficiency row.", {}),
    ("Topics: add, edit, delete, then try to delete Arrays to show the trigger message.", {}),
    ("Java and ODBC: run JdbcDemo --demo and OdbcDemo.ps1 -Demo; the same functions are called.", {}),
    ("Terminal: npm test, 31 of 31 passing.", {}),
], fill=LIGHT, pad=0.22, size=14)

# ---- 14. conclusion ---------------------------------------------------------------
s = slide("Conclusion", "Conclusion and Future Enhancements")
table(s, 0.667, 1.4, 12, [2.5, 9.5], ["Area", "Phase 3 outcome"], [
    ["Design to delivery", "Phase 1 EER and 3NF schema implemented in PostgreSQL and used by a working application"],
    ["Connectivity", "Web app through node-postgres; Java client through JDBC; PowerShell client through ODBC; all CRUD operations"],
    ["Database logic", "Constraints, functions, triggers and views enforce every rule for every client"],
    ["Quality", "23 documented test cases, 31 automated tests, all passing"],
    ["Beyond the plan", "pgvector similarity recommendations; indexes on foreign keys"],
], size=13, row_h=0.55)
text(s, 0.667, 4.9, 12, 1.8, [
    ("Next", {"bold": True, "size": 15, "color": NAVY}),
    ("Session tokens and role-based access · enforced quiz time limits · Redis leaderboard cache at scale "
     "· automatic recommendations after every attempt.", {"size": 14}),
], fill=LIGHT, pad=0.2, anchor=MSO_ANCHOR.MIDDLE)

prs.save(out)
print("saved", out, len(prs.slides), "slides")
